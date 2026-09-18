"""Automated Tests for Generic Excel Analysis, Aggregation, and Multi-Schema Workbook Intelligence."""

import os
import io
import pytest
import openpyxl
from decimal import Decimal
from starlette.testclient import TestClient

from backend.app.main import app
from backend.app.services.documents.excel_processor import excel_processor, format_inr, parse_decimal_safe, WorkbookType
from backend.app.services.conversation_context_manager import conversation_context_manager
from backend.app.services.routing.intent_router import intent_router, RouteRequest
from backend.app.schemas.company_ai import DepartmentEnum


@pytest.fixture
def upi_workbook_path():
    """Returns the path to the real uploaded UPI transaction workbook."""
    candidate_paths = [
        r"d:\MYGPT\data\uploads\d2468238-9020-4289-9cd8-b112277b2012.xlsx",
        r"d:\MYGPT\data\uploads\d5873dfb-d0b2-41d0-a850-7e18557e6b8a.xlsx",
    ]
    for p in candidate_paths:
        if os.path.exists(p):
            return p
    # If in different directory, search
    upload_dir = r"d:\MYGPT\data\uploads"
    for f in os.listdir(upload_dir):
        if f.endswith(".xlsx"):
            full = os.path.join(upload_dir, f)
            try:
                wb = openpyxl.load_workbook(full, read_only=True)
                if "Sheet1" in wb.sheetnames and wb["Sheet1"].max_row > 1000:
                    wb.close()
                    return full
                wb.close()
            except Exception:
                pass
    pytest.skip("UPI test workbook not found on local disk")


@pytest.fixture
def sprint_workbook_path():
    path = r"d:\MYGPT\sprint final.xlsx"
    if os.path.exists(path):
        return path
    pytest.skip("sprint final.xlsx not found")


class TestGenericExcelIntelligence:
    """Tests schema detection, numerical parsing, aggregation, and formatting."""

    def test_numeric_parsing(self):
        assert parse_decimal_safe("1000.0") == Decimal("1000.0")
        assert parse_decimal_safe("1,519,715.76") == Decimal("1519715.76")
        assert parse_decimal_safe("$2,250.00") == Decimal("2250.00")
        assert parse_decimal_safe("₹50,000.00") == Decimal("50000.00")
        assert parse_decimal_safe(None) is None
        assert parse_decimal_safe("NULL") is None
        assert parse_decimal_safe("-") is None

    def test_inr_formatting(self):
        assert format_inr(Decimal("1519715.76")) == "₹15,19,715.76"
        assert format_inr(Decimal("1087.06")) == "₹1,087.06"
        assert format_inr(Decimal("50000")) == "₹50,000.00"
        assert format_inr(Decimal("100")) == "₹100.00"

    def test_upi_workbook_schema_inspection(self, upi_workbook_path):
        inspection = excel_processor.inspect_workbook(upi_workbook_path)
        assert inspection["workbook_type"] == WorkbookType.TRANSACTION.value
        assert inspection["total_valid_rows"] == 1398
        assert "Sheet1" in inspection["sheets"]
        sheet_info = inspection["sheets"]["Sheet1"]
        assert sheet_info["modes"]["QR"] == 801
        assert sheet_info["modes"]["UPI"] == 593
        assert sheet_info["modes"]["BQR"] == 4
        assert sheet_info["statuses"]["AUTHORISED"] == 1398

    def test_sprint_workbook_schema_inspection(self, sprint_workbook_path):
        inspection = excel_processor.inspect_workbook(sprint_workbook_path)
        assert inspection["workbook_type"] == WorkbookType.SPRINT.value
        assert "Fazil" in inspection["owners"]
        assert "Reka" in inspection["owners"]

    def test_upi_aggregations(self, upi_workbook_path):
        # Total
        agg_total = excel_processor.calculate_aggregation(upi_workbook_path, target_column="Amount")
        assert agg_total["total"] == Decimal("1519715.76")
        assert agg_total["formatted_total"] == "₹15,19,715.76"
        assert agg_total["numeric_count"] == 1398

        # Average
        agg_avg = excel_processor.calculate_aggregation(upi_workbook_path, target_column="Amount", operation="AVG")
        assert round(float(agg_avg["average"]), 2) == 1087.06
        assert agg_avg["formatted_average"] == "₹1,087.06"

        # Settled Amount
        agg_settled = excel_processor.calculate_aggregation(upi_workbook_path, target_column="Settled_Amount")
        assert agg_settled["total"] == Decimal("1519715.76")

        # QR filter
        agg_qr = excel_processor.calculate_aggregation(upi_workbook_path, target_column="Amount", filter_criteria={"mode": "QR"})
        assert agg_qr["numeric_count"] == 801

        # UPI filter
        agg_upi = excel_processor.calculate_aggregation(upi_workbook_path, target_column="Amount", filter_criteria={"mode": "UPI"})
        assert agg_upi["numeric_count"] == 593

        # Min amount filter > 10000
        agg_high = excel_processor.calculate_aggregation(upi_workbook_path, target_column="Amount", filter_criteria={"min_amount": 10000})
        assert agg_high["numeric_count"] == 7


class TestEndToEndGenericExcelFlow:
    """Tests real API end-to-end queries against uploaded UPI workbook and Sprint workbook."""

    @pytest.fixture(autouse=True)
    def setup_client(self):
        self.client = TestClient(app)

    def test_e2e_upi_conversation_flow(self, upi_workbook_path):
        session_id = "test_upi_session_flow"
        conversation_context_manager.clear_state(session_id)

        with open(upi_workbook_path, "rb") as f:
            file_bytes = f.read()

        # 1. Upload UPI Excel
        upload_res = self.client.post(
            "/api/company/upload",
            files={"file": ("UPI - QR FROM 31-5-26 TO 15-06-26 1.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"session_id": session_id}
        )
        assert upload_res.status_code == 200
        upload_data = upload_res.json()
        assert upload_data["success"] is True
        assert upload_data["document_type"] == "EXCEL"
        assert upload_data["structured_data"]["workbook_type"] == "TRANSACTION"

        # 2. "give me a short summary"
        res_summary = self.client.post(
            "/api/company/chat",
            json={"message": "give me a short summary", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_summary.status_code == 200
        data_sum = res_summary.json()
        assert data_sum["status"] == "success"
        assert "1,398" in data_sum["answer"]
        assert "15,19,715.76" in data_sum["answer"]
        assert "QR" in data_sum["answer"]
        assert "UPI" in data_sum["answer"]
        assert "Total Sprint Tasks" not in data_sum["answer"]

        # 3. "what is the total amount in the excel"
        res_total = self.client.post(
            "/api/company/chat",
            json={"message": "what is the total amount in the excel", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_total.status_code == 200
        data_tot = res_total.json()
        assert data_tot["status"] == "success"
        assert "15,19,715.76" in data_tot["answer"]
        assert "GENERAL QUESTION" not in data_tot["answer"]

        # 4. "total amount" (short follow-up)
        res_short_tot = self.client.post(
            "/api/company/chat",
            json={"message": "total amount", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_short_tot.status_code == 200
        assert "15,19,715.76" in res_short_tot.json()["answer"]

        # 5. "how many transactions are there?"
        res_count = self.client.post(
            "/api/company/chat",
            json={"message": "how many transactions are there?", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_count.status_code == 200
        assert "1,398" in res_count.json()["answer"]

        # 6. "what is the total settled amount?"
        res_settled = self.client.post(
            "/api/company/chat",
            json={"message": "what is the total settled amount?", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_settled.status_code == 200
        assert "15,19,715.76" in res_settled.json()["answer"]

        # 7. "how many QR transactions?"
        res_qr = self.client.post(
            "/api/company/chat",
            json={"message": "how many QR transactions?", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_qr.status_code == 200
        assert "801" in res_qr.json()["answer"]

        # 8. "how many UPI transactions?"
        res_upi = self.client.post(
            "/api/company/chat",
            json={"message": "how many UPI transactions?", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_upi.status_code == 200
        assert "593" in res_upi.json()["answer"]

        # 9. "what is the average transaction amount?"
        res_avg = self.client.post(
            "/api/company/chat",
            json={"message": "what is the average transaction amount?", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_avg.status_code == 200
        assert "1,087.06" in res_avg.json()["answer"]

        # 10. "show me transactions above 10000"
        res_above = self.client.post(
            "/api/company/chat",
            json={"message": "show me transactions above 10000", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_above.status_code == 200
        assert "7" in res_above.json()["answer"]
        assert "60,000.00" in res_above.json()["answer"]

        # 11. "give me a report"
        res_rep = self.client.post(
            "/api/company/chat",
            json={"message": "give me a report", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert res_rep.status_code == 200
        assert "15,19,715.76" in res_rep.json()["answer"]
        assert "TRANSACTION AUDIT REPORT" in res_rep.json()["answer"]

    def test_sprint_split_backward_compatibility(self, sprint_workbook_path):
        session_id = "test_sprint_compatibility"
        conversation_context_manager.clear_state(session_id)

        with open(sprint_workbook_path, "rb") as f:
            file_bytes = f.read()

        upload_res = self.client.post(
            "/api/company/upload",
            files={"file": ("sprint final.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"session_id": session_id}
        )
        assert upload_res.status_code == 200

        chat_res = self.client.post(
            "/api/company/chat",
            json={"message": "create separate sprint sheet for Fazil and Reka", "session_id": session_id, "user_role": "EMPLOYEE"}
        )
        assert chat_res.status_code == 200
        chat_data = chat_res.json()
        assert chat_data["status"] == "success"
        files = chat_data["files"]
        assert len(files) >= 2
        file_names = [f["filename"] for f in files]
        assert "Fazil_Sprint.xlsx" in file_names
        assert "Reka_Sprint.xlsx" in file_names
